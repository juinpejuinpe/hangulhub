"""Parse uploaded documents into structured table/text content."""

import csv
import io
import re
from pathlib import Path

from store import clean_text

MAX_CELL = 2000
MAX_ROWS = 20000


def parse_file(filename: str, data: bytes) -> dict:
    ext = Path(filename).suffix.lower().lstrip(".")
    name = Path(filename).name
    if ext in ("xlsx", "xlsm"):
        return _parse_xlsx(data)
    if ext == "xls":
        raise ValueError("Old .xls files are not supported yet. Please save as .xlsx.")
    if ext in ("csv", "tsv"):
        return _parse_delimited(data, ext)
    if ext in ("txt", "md"):
        return _parse_text(data)
    if ext == "docx":
        return _parse_docx(data)
    if ext == "pdf":
        return _parse_pdf(data)
    raise ValueError(
        f"Unsupported file type .{ext}. Supported: .xlsx .csv .tsv .txt .md .docx .pdf"
    )


def _truncate(value):
    value = clean_text(value)
    return value[:MAX_CELL]


def _parse_xlsx(data: bytes) -> dict:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheet = wb.worksheets[0]
    rows = []
    for row in sheet.iter_rows(values_only=True):
        rows.append([_truncate(c) for c in row])
        if len(rows) >= MAX_ROWS:
            break
    wb.close()
    return _table_from_rows(rows)


def _parse_delimited(data: bytes, ext: str) -> dict:
    delim = "\t" if ext == "tsv" else ","
    text = data.decode("utf-8-sig", errors="replace")
    reader = csv.reader(io.StringIO(text), delimiter=delim)
    rows = []
    for row in reader:
        row = [_truncate(c) for c in row]
        if any(cell.strip() for cell in row):
            rows.append(row)
        if len(rows) >= MAX_ROWS:
            break
    return _table_from_rows(rows)


def _table_from_rows(rows: list) -> dict:
    if not rows:
        return {"kind": "table", "headers": [], "rows": [], "note": "Empty file."}
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    first = rows[0]
    has_header = (
        any(str(c).strip() for c in first)
        and not all(_looks_numeric(c) for c in first)
        and len(rows) > 1
        and any(str(rows[0][i]).strip() != str(rows[1][i]).strip()
                for i in range(width))
    )
    if has_header:
        headers = [str(c).strip() or f"Column {i + 1}" for i, c in enumerate(first)]
        data_rows = rows[1:]
    else:
        headers = [f"Column {i + 1}" for i in range(width)]
        data_rows = rows
    dict_rows = []
    for r in data_rows:
        row = {}
        for i, h in enumerate(headers):
            row[h] = r[i] if i < len(r) else ""
        if any(row[h].strip() for h in headers):
            dict_rows.append(row)
    return {
        "kind": "table",
        "headers": headers,
        "rows": dict_rows,
        "headerRow": has_header,
    }


def _looks_numeric(value) -> bool:
    try:
        float(str(value).replace(",", "").strip())
        return True
    except ValueError:
        return False


def _parse_text(data: bytes) -> dict:
    text = data.decode("utf-8-sig", errors="replace")
    lines = [line.strip() for line in text.splitlines()]
    lines = [line for line in lines if line]
    return {
        "kind": "text",
        "text": "\n".join(lines),
        "lines": lines[:MAX_ROWS],
    }


def _parse_docx(data: bytes) -> dict:
    import docx

    document = docx.Document(io.BytesIO(data))
    chunks = []
    for para in document.paragraphs:
        if para.text.strip():
            chunks.append(clean_text(para.text))
    for table in document.tables:
        for row in table.rows:
            cells = [clean_text(cell.text) for cell in row.cells]
            if any(cells):
                chunks.append("\t".join(cells))
    text = "\n".join(chunks)
    lines = [line.strip() for line in text.splitlines()]
    return {
        "kind": "text",
        "text": "\n".join(lines),
        "lines": [line for line in lines if line][:MAX_ROWS],
    }


def _parse_pdf(data: bytes) -> dict:
    import pdfplumber

    chunks = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text() or ""
            if page_text.strip():
                chunks.append(page_text)
    text = "\n\n".join(chunks)
    if not text.strip():
        raise ValueError(
            "No extractable text found in this PDF. It may be a scanned image; "
            "please provide a digital PDF."
        )
    return {
        "kind": "text",
        "text": text,
        "lines": text.splitlines()[:MAX_ROWS],
        "pages": len(chunks),
    }


def rows_to_pairs(content: dict, front_col: str, back_cols: list) -> list:
    """Build (front, back) pairs from a table file."""
    pairs = []
    headers = content.get("headers", [])
    for row in content.get("rows", []):
        front = clean_text(row.get(front_col, ""))
        if not front:
            continue
        backs = [clean_text(row.get(h, "")) for h in back_cols if h in headers]
        back = "\n".join(b for b in backs if b) if backs else ""
        pairs.append((front, back))
    return pairs


def text_to_pairs(content: dict) -> list:
    """Best-effort line pairing for plain-text vocab lists."""
    pairs = []
    delimiters = re.compile(r"\t|\s{2,}| \| | [-=] |::?|: ")
    hangul = re.compile(r"[\uac00-\ud7af\u1100-\u11ff\u3130-\u318f]")
    latin = re.compile(r"[A-Za-z]")
    for line in content.get("lines", []):
        if not line:
            continue
        parts = re.split(delimiters, line, maxsplit=1)
        parts = [p.strip() for p in parts]
        if len(parts) >= 2 and parts[0] and parts[1]:
            pairs.append((parts[0], parts[1]))
            continue
        # PDF/text lines often look like "학교 school": Korean first, then
        # English. Split at the first Latin character that follows Hangul.
        first_hangul = hangul.search(line)
        if first_hangul:
            boundary = None
            for idx, ch in enumerate(line):
                if idx > first_hangul.start() and latin.match(ch):
                    boundary = idx
                    break
            if boundary:
                front = line[:boundary].strip()
                back = line[boundary:].strip()
                if front and back:
                    pairs.append((front, back))
    return pairs
